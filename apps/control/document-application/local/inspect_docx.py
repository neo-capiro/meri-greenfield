"""Fixed fixture-only package/content check. No Office renderer or network use."""
import hashlib
import io
import json
import sys
import zipfile


def main():
    if sys.argv[1:] == ['--help']:
        print('Read one bounded DOCX from stdin; report SHA256, paragraphs and tables. Exit 0 success/help; 2 failure.')
        return 0
    if sys.argv[1:]:
        return 2
    raw = sys.stdin.buffer.read(1048577)
    if not raw or len(raw) > 1048576:
        return 2
    from docx import Document
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if len(archive.infolist()) > 5000 or sum(entry.file_size for entry in archive.infolist()) > 33554432 or archive.testzip() is not None:
            return 2
    document = Document(io.BytesIO(raw))
    print(json.dumps({'sha256': hashlib.sha256(raw).hexdigest(), 'bytes': len(raw),
                     'paragraphs': [paragraph.text for paragraph in document.paragraphs],
                     'tables': [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
                     'rendered': False, 'nativeQualified': False}))
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception:
        print('DOCX_FIXTURE_INSPECTION_FAILED', file=sys.stderr)
        sys.exit(2)
