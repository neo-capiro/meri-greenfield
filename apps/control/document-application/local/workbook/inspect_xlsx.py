"""Bounded local workbook inspection using openpyxl; no generator or formula evaluator."""
import hashlib
import io
import json
import sys
import zipfile


def main():
    if sys.argv[1:] == ['--help']:
        print('Read a bounded XLSX from stdin; inspect cells, formulas, cached values and package members. Exit0 success/help;2 failure.')
        return 0
    if sys.argv[1:]:
        return 2
    raw = sys.stdin.buffer.read(1048577)
    if not raw or len(raw) > 1048576:
        raise ValueError('Workbook size')
    import openpyxl
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        members = archive.namelist()
        if len(members) > 5000 or sum(x.file_size for x in archive.infolist()) > 33554432 or archive.testzip() is not None:
            raise ValueError('Workbook package')
        if any('vbaProject' in x or x.startswith('xl/externalLinks/') for x in members):
            raise ValueError('Active or linked workbook')
    formulas = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
    try:
        if len(formulas.sheetnames) > 10:
            raise ValueError('Sheet limit')
        sheets = []
        for sheet in formulas:
            if sheet.max_row * sheet.max_column > 20000:
                raise ValueError('Cell limit')
            cells = [{'address': cell.coordinate, 'value': cell.value, 'kind': cell.data_type,
                      'numberFormat': cell.number_format, 'cached': cached[sheet.title][cell.coordinate].value}
                     for row in sheet for cell in row if cell.value is not None]
            sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
                           'cells': cells, 'charts': len(sheet._charts)})
        print(json.dumps({'sha256': hashlib.sha256(raw).hexdigest(), 'bytes': len(raw), 'sheets': sheets,
                          'members': members, 'nativeQualified': False}, allow_nan=False))
    finally:
        formulas.close()
        cached.close()
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception:
        print('XLSX_FIXTURE_INSPECTION_FAILED', file=sys.stderr)
        sys.exit(2)
