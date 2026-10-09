#!/usr/bin/env python3
"""Meri W02-01 synthetic Office generation/inspection probe (opt-in).

Generates native DOCX (heading + table) and XLSX (supported formula + known
values) with the adopted Python libraries, reopens them with the same
libraries, inspects the actual OOXML/ZIP structure and inspects PDFs with
pypdf. Deterministic fixture expectations; no visual inspection of a
nonexistent output. Rendering (LibreOffice/poppler) and XLSX recalculation are
driven by probes/artifacts.mjs and are live gates, not performed here.

Never invoked by `npm test`. Uses only scoped installed dependencies and the
owned output directory passed on the command line. Every nested failure is
propagated to the top-level `ok`, and observed values are reported SEPARATELY
from verdicts so a wrong value, a failed command or a parse error never
collapses into a plausible-looking success.
"""

import argparse
import hashlib
import json
import os
import re
import sys
import zipfile

DOCX_HEADING = "Meri qualification heading"
DOCX_TABLE = [["label", "value"], ["answer", "6"]]
XLSX_INPUTS = [1, 2, 3]
XLSX_FORMULA = "=SUM(A1:A3)"
XLSX_EXPECTED = 6


def sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def zip_structure(path: str) -> list:
    with zipfile.ZipFile(path) as zf:
        return sorted(zf.namelist())


def _zip_read_bounded(path: str, member: str, limit: int = 600) -> str:
    try:
        with zipfile.ZipFile(path) as zf:
            return zf.read(member).decode("utf-8", "replace")[:limit]
    except Exception:
        return ""


def _xlsx_facts(path: str) -> dict:
    sheet = _zip_read_bounded(path, "xl/worksheets/sheet1.xml", 4000)
    m = re.search(r"<c r=\"A4\".*?</c>", sheet, re.S)
    calc = re.search(r"<calcPr[^>]*/>", _zip_read_bounded(path, "xl/workbook.xml", 4000))
    return {
        "a4Xml": (m.group(0) if m else "")[:400],
        "calcProperties": (calc.group(0) if calc else ""),
    }


def generate_docx(outdir: str) -> dict:
    import docx  # python-docx

    out = os.path.join(outdir, "qualification.docx")
    document = docx.Document()
    document.add_heading(DOCX_HEADING, level=1)
    document.add_paragraph("Synthetic body paragraph for native DOCX qualification.")
    table = document.add_table(rows=2, cols=2)
    for r, row in enumerate(DOCX_TABLE):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.save(out)

    reopened = docx.Document(out)
    headings = [p.text for p in reopened.paragraphs if p.style.name.startswith("Heading")]
    table_text = [[cell.text for cell in row.cells] for row in reopened.tables[0].rows] if reopened.tables else []
    structure = zip_structure(out)
    ok = (headings == [DOCX_HEADING] and len(reopened.tables) == 1
          and table_text == DOCX_TABLE and "word/document.xml" in structure)
    return {
        "file": out, "sha256": sha256(out), "ok": ok,
        "expected": {"heading": DOCX_HEADING, "table": DOCX_TABLE},
        "reopened": {"headings": headings, "tableText": table_text, "tableCount": len(reopened.tables)},
        "structure": structure, "hasDocumentXml": "word/document.xml" in structure,
    }


def generate_xlsx(outdir: str) -> dict:
    import xlsxwriter  # XlsxWriter

    out = os.path.join(outdir, "qualification.xlsx")
    workbook = xlsxwriter.Workbook(out)
    sheet = workbook.add_worksheet("qual")
    for row, value in enumerate(XLSX_INPUTS):
        sheet.write_number(row, 0, value)
    # Documented XlsxWriter blank cache so non-Excel apps recalculate rather than
    # show the default cached zero. Never seed the expected result here.
    # https://xlsxwriter.readthedocs.io/faq.html#q-why-do-my-formulas-show-a-zero-result-in-some-non-excel-applications
    sheet.write_formula(3, 0, XLSX_FORMULA, None, "")
    workbook.close()

    import openpyxl

    ws = openpyxl.load_workbook(out, data_only=False).active
    cached = openpyxl.load_workbook(out, data_only=True).active["A4"].value
    structure = zip_structure(out)
    inputs = [ws.cell(row=r, column=1).value for r in range(1, 4)]
    facts = _xlsx_facts(out)
    ok = (inputs == XLSX_INPUTS and ws["A4"].value == XLSX_FORMULA
          and "xl/worksheets/sheet1.xml" in structure)
    return {
        "file": out, "sha256": sha256(out), "ok": ok,
        "expected": {"inputs": XLSX_INPUTS, "formula": XLSX_FORMULA, "recalculatedResult": XLSX_EXPECTED},
        "reopened": {"inputs": inputs, "a4Formula": ws["A4"].value, "a4GeneratedCachedValue": cached},
        "cachedValueNote": "the generated fixture uses a blank cached result; a cached value that is absent/zero here does not prove LibreOffice recalculation or Excel fidelity",
        "structure": structure, "hasSheetXml": "xl/worksheets/sheet1.xml" in structure,
        "xmlFacts": facts,
    }


def inspect_file(path: str) -> dict:
    structure = zip_structure(path)
    return {"file": path, "sha256": sha256(path), "structure": structure, "ok": len(structure) > 0}


def inspect_pptx(path: str) -> dict:
    structure = zip_structure(path)
    slides = [n for n in structure if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    notes = [n for n in structure if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")]
    ok = len(slides) == 2 and len(notes) >= 2
    return {
        "file": path, "sha256": sha256(path), "ok": ok,
        "slideCount": len(slides), "notesCount": len(notes), "structure": structure,
    }


def inspect_pdf(path: str) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(path)
    text = ""
    for page in reader.pages[:3]:
        text += (page.extract_text() or "")
    return {
        "file": path, "sha256": sha256(path), "ok": len(reader.pages) >= 1 and len(text.strip()) > 0,
        "pageCount": len(reader.pages), "firstPageTextSample": text[:500],
    }


def xlsx_recalc(path: str) -> dict:
    """Reopen a LibreOffice-recalculated XLSX copy and report observed facts.

    Requires original inputs (1,2,3), a retained `=SUM(A1:A3)` formula and a
    NUMERIC cached result of 6. A hand-typed constant 6 with no formula fails.
    """
    import openpyxl

    wb_f = openpyxl.load_workbook(path, data_only=False)
    ws_f = wb_f.active
    inputs = [ws_f.cell(row=r, column=1).value for r in range(1, 4)]
    formula = ws_f["A4"].value
    cached = openpyxl.load_workbook(path, data_only=True).active["A4"].value
    observed = {
        "inspected": True,
        "inputs": inputs,
        "formula": formula,
        "cachedValue": cached,
        "cachedType": type(cached).__name__,
    }
    ok = (inputs == XLSX_INPUTS and formula == XLSX_FORMULA
          and isinstance(cached, (int, float)) and not isinstance(cached, bool) and cached == XLSX_EXPECTED)
    return {
        "file": path, "sha256": sha256(path), "ok": ok,
        "expected": {"inputs": XLSX_INPUTS, "formula": XLSX_FORMULA, "recalculatedResult": XLSX_EXPECTED},
        "observed": observed,
        "observedA4": cached,
        "xmlFacts": _xlsx_facts(path),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Meri qualification Office artifact probe")
    sub = parser.add_subparsers(dest="command", required=True)
    gen = sub.add_parser("generate"); gen.add_argument("outdir")
    insp = sub.add_parser("inspect"); insp.add_argument("path")
    pdfc = sub.add_parser("pdf"); pdfc.add_argument("path")
    rc = sub.add_parser("xlsxrecalc"); rc.add_argument("path")
    args = parser.parse_args()

    if args.command == "generate":
        os.makedirs(args.outdir, exist_ok=True)
        results = [generate_docx(args.outdir), generate_xlsx(args.outdir)]
    elif args.command == "inspect":
        if args.path.endswith(".pptx"):
            results = [inspect_pptx(args.path)]
        else:
            results = [inspect_file(args.path)]
    elif args.command == "xlsxrecalc":
        results = [xlsx_recalc(args.path)]
    else:
        results = [inspect_pdf(args.path)]

    payload = {"probe": "artifacts-python", "ok": all(r["ok"] for r in results), "results": results}
    print(json.dumps(payload, indent=2))
    return 0 if payload["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
