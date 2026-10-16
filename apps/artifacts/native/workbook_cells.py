"""Bounded workbook cell reader for native checks: every non-empty cell's formula text and stored (cached) value.

Reads one XLSX from stdin and prints JSON {sheets: {name: {address: {formula, value}}}}. It evaluates nothing: `value` is
whatever the file stores, so it shows a recalculation only after a real office suite recalculated and saved the file.
Exit 0 success, 2 failure (message on stderr carries no workbook content).
"""
import io
import json
import sys
import zipfile

MAX_BYTES = 8388608
MAX_CELLS = 20000


def main():
    if sys.argv[1:] == ['--help']:
        print(__doc__.strip())
        return 0
    if sys.argv[1:]:
        return 2
    raw = sys.stdin.buffer.read(MAX_BYTES + 1)
    if not raw or len(raw) > MAX_BYTES:
        raise ValueError('size')
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if any('vbaProject' in name or name.startswith('xl/externalLinks/') for name in archive.namelist()):
            raise ValueError('active or linked workbook')
    import openpyxl
    formulas = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    stored = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
    try:
        sheets, count = {}, 0
        for sheet in formulas:
            cells = {}
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    count += 1
                    if count > MAX_CELLS:
                        raise ValueError('cells')
                    text = cell.value if cell.data_type == 'f' and isinstance(cell.value, str) else None
                    cells[cell.coordinate] = {'formula': text, 'value': stored[sheet.title][cell.coordinate].value}
            sheets[sheet.title] = cells
        print(json.dumps({'sheets': sheets}, allow_nan=False, default=str))
    finally:
        formulas.close()
        stored.close()
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as error:  # closed failure; no workbook content
        print('WORKBOOK_CELLS_FAILED ' + type(error).__name__, file=sys.stderr)
        sys.exit(2)
