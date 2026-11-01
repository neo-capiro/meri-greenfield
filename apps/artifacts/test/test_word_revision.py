"""Native retained-base revision properties, without network or source access."""
import importlib.util
import io
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

spec = importlib.util.spec_from_file_location('meri_office', Path(__file__).parents[1] / 'python' / 'office.py')
office = importlib.util.module_from_spec(spec)
spec.loader.exec_module(office)


class WordRevision(unittest.TestCase):
    def test_quote_depth_and_visible_marks_survive_native_docx_reopen(self):
        blocks = [{'kind': 'paragraph', 'quoteDepth': depth, 'runs': [
            {'text': 'Deleted', 'strike': True, 'bold': True},
            {'text': 'Inserted', 'underline': True, 'italic': True},
            {'text': 'Explicit off', 'strike': False, 'underline': False},
            {'text': 'Unmarked'},
        ]} for depth in (1, 8)]
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'structured.docx'
            office.word({'blocks': blocks}, {}, output)
            document = Document(output)
            for paragraph, twips in zip(document.paragraphs, (360, 2880)):
                ind = paragraph._p.pPr.find(qn('w:ind'))
                self.assertEqual(ind.get(qn('w:left')), str(twips))
                self.assertEqual(ind.get(qn('w:right')), str(twips))
                deleted, inserted, explicit_off, unmarked = paragraph.runs
                self.assertTrue(deleted.font.strike)
                self.assertTrue(deleted.bold)
                self.assertTrue(inserted.underline)
                self.assertTrue(inserted.italic)
                self.assertFalse(explicit_off.font.strike)
                self.assertFalse(explicit_off.underline)
                self.assertIsNone(unmarked.font.strike)
                self.assertIsNone(unmarked.underline)
            self.assertEqual(len(document.element.xpath('.//w:ins | .//w:del')), 0)

    def test_revision_preserves_inherited_formatting_when_new_fields_are_omitted(self):
        original = Document()
        normal = original.styles['Normal']
        normal.paragraph_format.left_indent = Inches(0.4)
        normal.paragraph_format.right_indent = Inches(0.6)
        normal.font.strike = True
        normal.font.underline = True
        original.sections[0].right_margin = Inches(1.4)
        original.add_paragraph('Immutable parent')
        raw = io.BytesIO()
        original.save(raw)
        before = raw.getvalue()
        blocks = [
            {'kind': 'paragraph', 'runs': [{'text': 'Inherited'}]},
            {'kind': 'paragraph', 'quoteDepth': 2, 'runs': [
                {'text': 'Override', 'strike': False, 'underline': False}]},
        ]
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'revision.docx'
            office.word({'blocks': blocks}, {}, output, {'bytes': before})
            revised = Document(output)
            inherited, overridden = revised.paragraphs
            self.assertIsNone(inherited.paragraph_format.left_indent)
            self.assertIsNone(inherited.paragraph_format.right_indent)
            self.assertIsNone(inherited.runs[0].font.strike)
            self.assertIsNone(inherited.runs[0].underline)
            self.assertEqual(revised.styles['Normal'].element.xml, normal.element.xml)
            self.assertEqual(revised.sections[0]._sectPr.xml, original.sections[0]._sectPr.xml)
            self.assertEqual(overridden.paragraph_format.left_indent, Inches(0.5))
            self.assertEqual(overridden.paragraph_format.right_indent, Inches(0.5))
            self.assertFalse(overridden.runs[0].font.strike)
            self.assertFalse(overridden.runs[0].underline)
        self.assertEqual(raw.getvalue(), before)
        self.assertEqual(Document(io.BytesIO(before)).paragraphs[0].text, 'Immutable parent')

    def test_native_base_styles_sections_and_original_survive_body_replacement(self):
        original = Document()
        original.sections[0].left_margin = Inches(1.3)
        original.styles['Normal'].font.name = 'Liberation Serif'
        original.add_paragraph('Original text')
        original.add_table(rows=1, cols=1).cell(0, 0).text = 'Original table'
        raw = io.BytesIO()
        original.save(raw)
        before = raw.getvalue()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'revision.docx'
            result = office.word({'blocks': [{'kind': 'paragraph', 'runs': [{'text': 'Revised text'}]}]}, {}, output,
                                {'request': {'format': 'docx'}, 'bytes': before})
            revised = Document(output)
            self.assertEqual([p.text for p in revised.paragraphs], ['Revised text'])
            self.assertEqual(len(revised.tables), 0)
            self.assertEqual(revised.sections[0].left_margin, Inches(1.3))
            self.assertEqual(revised.styles['Normal'].font.name, 'Liberation Serif')
            self.assertEqual(result, {'paragraphs': 1, 'tables': 0, 'pictures': 0})
        reopened = Document(io.BytesIO(before))
        self.assertEqual(reopened.paragraphs[0].text, 'Original text')
        self.assertEqual(reopened.tables[0].cell(0, 0).text, 'Original table')
        self.assertEqual(raw.getvalue(), before)

    def test_invalid_selected_base_cannot_be_ignored(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(Exception):
                office.word({'blocks': [{'kind': 'paragraph', 'runs': [{'text': 'Replacement'}]}]}, {},
                            Path(directory) / 'refused.docx', {'bytes': b'not a DOCX'})
            self.assertFalse((Path(directory) / 'refused.docx').exists())


if __name__ == '__main__':
    unittest.main()
