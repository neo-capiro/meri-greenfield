"""Original data-only Office builders over adopted upstream libraries."""
import io

def _excel_column_widths(rows):
    """Bounded default-font estimates, not formula evaluation or universal autofit."""
    from math import ceil
    from xlsxwriter.utility import cell_autofit_width
    widths={}
    for row in rows:
        for index,cell in enumerate(row):
            value=cell.get('value')
            if 'formula' in cell:
                # A result allowance: the expression itself is never displayed.
                text='0'*14
            elif isinstance(value,bool):text='TRUE' if value else 'FALSE'
            else:text='' if value is None else str(value)
            pixels=max((cell_autofit_width(line) for line in text.splitlines()),default=0)
            if cell.get('numberFormat') and ('formula' in cell or isinstance(value,(int,float))):
                # Room for common separators/currency without a second formatter.
                pixels=max(pixels,cell_autofit_width(cell['numberFormat'])+28)
            widths[index]=max(widths.get(index,12),min(48,max(12,ceil(pixels/7)+2)))
    return widths

def word(spec,assets,path):
    from docx import Document
    from docx.shared import Inches
    document=Document()
    for block in spec['blocks']:
        if block['kind']=='paragraph':
            p=document.add_paragraph(style=block.get('style','Normal'))
            for item in block['runs']:
                r=p.add_run(item['text']);r.bold=item.get('bold',False);r.italic=item.get('italic',False)
        elif block['kind']=='table':
            table=document.add_table(rows=len(block['rows']),cols=len(block['rows'][0]));table.style='Table Grid'
            for row,values in zip(table.rows,block['rows']):
                for cell,value in zip(row.cells,values):cell.text=value
        elif block['kind']=='image':document.add_picture(io.BytesIO(assets[block['assetId']]),width=Inches(block['widthInches']),height=Inches(block['heightInches']))
        elif block['kind']=='chart':
            document.add_picture(io.BytesIO(chart_png(block['chart'])),width=Inches(6))
        else:raise ValueError('Unsupported Word block')
    document.save(path)
    # Reopening checks package readability, not Office rendering or visual quality.
    opened=Document(path)
    return {'paragraphs':len(opened.paragraphs),'tables':len(opened.tables),'pictures':len(opened.inline_shapes)}

def chart_png(chart):
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt
    fig,ax=plt.subplots(figsize=(8,4))
    try:
        for series in chart['series']:
            if chart['kind']=='line':ax.plot(chart['categories'],series['values'],label=series['name'])
            else:ax.bar(chart['categories'],series['values'],label=series['name'],alpha=.7)
        ax.set_title(chart['title']);ax.legend();fig.tight_layout();out=io.BytesIO();fig.savefig(out,format='png',dpi=100);return out.getvalue()
    finally:plt.close(fig)

def excel(spec,assets,path,base=None):
    import xlsxwriter,openpyxl
    if base and all(not s.get('charts') and not s.get('images') for s in [*base['request']['content']['sheets'],*spec['sheets']]) and [s['name'] for s in base['request']['content']['sheets']]==[s['name'] for s in spec['sheets']]:
        # Only our retained, typed plain workbook source can enter this edit path.
        # Arbitrary imported workbooks/macros/links are not admitted.
        book=openpyxl.load_workbook(io.BytesIO(base['bytes']),data_only=False,keep_links=False)
        try:
            for sheet in spec['sheets']:
                page=book[sheet['name']]
                for row in page:
                    for cell in row:cell.value=None;cell.number_format='General'
                for y,row in enumerate(sheet['rows'],1):
                    for x,item in enumerate(row,1):
                        cell=page.cell(y,x)
                        if 'formula' in item:cell.value=item['formula']
                        else:
                            cell.value=item.get('value')
                            if isinstance(cell.value,str):cell.data_type='s'
                        if item.get('numberFormat'):cell.number_format=item['numberFormat']
                for index,width in _excel_column_widths(sheet['rows']).items():
                    column=page.column_dimensions[openpyxl.utils.get_column_letter(index+1)]
                    # Split any grouped dimensions produced by XlsxWriter.
                    column.min=column.max=index+1;column.width=width
            book.calculation.fullCalcOnLoad=True;book.calculation.forceFullCalc=True;book.save(path)
        finally:book.close()
        return {'sheets':[s['name'] for s in spec['sheets']],'formulas':sum(1 for s in spec['sheets'] for row in s['rows'] for cell in row if 'formula' in cell),'formulaCache':'removed/unverified; recalculation required','revisionMethod':'openpyxl typed-cell replacement'}
    with xlsxwriter.Workbook(path,{'strings_to_urls':False,'strings_to_formulas':False,'strings_to_numbers':False}) as book:
        for sheet in spec['sheets']:
            page=book.add_worksheet(sheet['name'])
            for index,width in _excel_column_widths(sheet['rows']).items():
                page.set_column(index,index,width)
            for y,row in enumerate(sheet['rows']):
                for x,cell in enumerate(row):
                    fmt=book.add_format({'num_format':cell['numberFormat']}) if cell.get('numberFormat') else None
                    if 'formula' in cell:page.write_formula(y,x,cell['formula'],fmt,0)
                    elif isinstance(cell.get('value'),str):page.write_string(y,x,cell['value'],fmt)
                    elif cell.get('value') is not None:page.write(y,x,cell['value'],fmt)
            chart_row=len(sheet['rows'])+2
            for chart in sheet.get('charts',[]):
                output=book.add_chart({'type':'column' if chart['kind']=='bar' else 'line'});output.set_title({'name':chart['title']})
                for i,cat in enumerate(chart['categories']):page.write_string(chart_row+i,0,cat)
                for j,series in enumerate(chart['series']):
                    for i,val in enumerate(series['values']):page.write_number(chart_row+i,j+1,val)
                    output.add_series({'name':series['name'],'categories':[sheet['name'],chart_row,0,chart_row+len(chart['categories'])-1,0],'values':[sheet['name'],chart_row,j+1,chart_row+len(chart['categories'])-1,j+1]})
                page.insert_chart(chart_row,12,output);chart_row+=max(len(chart['categories'])+2,16)
            for i,image in enumerate(sheet.get('images',[])):
                from PIL import Image
                with Image.open(io.BytesIO(assets[image['assetId']])) as source:width,height=source.size
                page.insert_image(i*15,26,'image.png',{'image_data':io.BytesIO(assets[image['assetId']]),'x_scale':image['widthInches']*96/width,'y_scale':image['heightInches']*96/height})
    book=openpyxl.load_workbook(path,read_only=True,data_only=False)
    try:return {'sheets':book.sheetnames,'formulas':sum(1 for page in book for row in page for cell in row if cell.data_type=='f'),'formulaCache':'unverified-zero-placeholder; recalculation required'}
    finally:book.close()
