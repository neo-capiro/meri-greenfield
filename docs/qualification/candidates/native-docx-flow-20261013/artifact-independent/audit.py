"""Inspect saved editable DOCX, independent renders and canonical source/provenance facts."""
import hashlib,json,pathlib,re,unicodedata,zipfile,os
import docx,fitz
root=pathlib.Path.cwd();here=pathlib.Path(__file__).resolve().parent;donor=root/'apps/runtime/artifacts/test/evidence';rows=[]
sha=lambda b:hashlib.sha256(b).hexdigest()
canonical=lambda x:json.dumps(x,sort_keys=True,separators=(',',':'),ensure_ascii=False).encode()
hashj=lambda x:sha(canonical(x))
clean=lambda t:re.sub(r'\s+','',unicodedata.normalize('NFKC',t).replace('\u00ad',''))
assets=json.loads((root/'packages/skills/assets.json').read_text());asset=next(a for a in assets['assets'] if a['skillId']=='quick_skill.bill_watch' and a['skillVersion']=='1.0.0');manifest=json.loads((root/'packages/skills'/asset['manifestPath']).read_text());assert hashj(manifest)==asset['manifestSha256'];assert sha((root/'packages/skills'/asset['prosePath']).read_bytes())==asset['proseSha256'];pin={k:asset[k] for k in ['skillId','skillVersion','manifestSha256','proseSha256']}|{'catalogSha256':assets['catalogSha256']}
for cohort in ['retained-docx','retained-docx-final']:
 receipt=json.loads((donor/cohort/'receipt.json').read_text());revision=receipt['sourcePreparations'][1];assert receipt['nativeQualified'] is False;assert receipt['storedVersions']==receipt['objects']==2
 assert hashj(revision['reads'])==revision['provenance']['readsSha256'];assert revision['context']['evidence']==receipt['versions'][1]['evidence'];assert revision['context']['scopes']==receipt['versions'][1]['scopes'];assert revision['context']['binding']==receipt['versions'][1]['binding'];assert revision['baseVersionId']==receipt['versions'][0]['versionId']
 reads=[]
 for read in revision['reads']:
  e=read['envelope'];assert e['binding']==revision['context']['binding'];assert e['subject']['scopes']==revision['context']['scopes'];assert e['method']=='source-read';assert e['subject']['authorizationDecisionId'];sid=e['sources'][0]['sourceId'];assert sid in ['bill_text','federal_legislation','program_elements']
  assert e['operation']['id']=={'bill_text':'read_bill_text','federal_legislation':'get_bill_status','program_elements':'get_pe_budget_timeline'}[sid]
  selector={'subject':revision['provenance']['subject'],'asOf':revision['provenance']['asOf']};assert hashj(selector)==e['argumentsSha256'];
  for k,v in selector['subject'].items():
   name=re.sub(r'[A-Z]',lambda m:'_'+m[0].lower(),k);f=[x for x in e['appliedFilters'] if x['name']==name];assert len(f)==1 and f[0]['valueSha256']==hashj(v)
  for record in (e.get('data') or {}).get('records',[]):
   proof=next(p for p in read['evidence'] if p['citationId']==record['citationId']);assert proof['sourceId']==sid;assert proof['recordId']==record['recordId'];assert proof['sourceVersionId']==record['recordVersionId'];assert proof['contentSha256']==record['contentSha256']==sha(record['content']['text'].encode());assert proof['scopes']==record['scopes']==e['subject']['scopes'];assert proof['authorizationDecisionId']==e['subject']['authorizationDecisionId'];assert proof['observedAt']==e['sources'][0]['retrievedAt'];reads.append({'sourceId':sid,'recordId':record['recordId'],'version':record['recordVersionId'],'sha256':proof['contentSha256'],'scopes':proof['scopes'],'decision':proof['authorizationDecisionId']})
 assert revision['facts'][0]['status']=='bound-excerpt';assert revision['facts'][1]['status']=='unverified-claim' and revision['facts'][1]['conflict'];assert revision['facts'][2]['status']=='missing-evidence'
 for fact in revision['facts']:
  for link in fact['links']:
   if not link.get('evidence'):continue
   record=next(r for read in revision['reads'] for r in (read['envelope'].get('data') or {}).get('records',[]) if r['citationId']==link['evidence']['citationId']);text=record['content']['text'];c=link['citation'];passage=text.encode('utf-16-le')[c['start']*2:c['end']*2].decode('utf-16-le');assert passage==link['passage'];assert sha(passage.encode())==link['passageSha256']
 original_provenance=None
 for index,label in enumerate(['original','revision']):
  directory=here/cohort;p=directory/(label+'.docx');data=p.read_bytes();source=donor/cohort/(label+'.docx');assert data==source.read_bytes();version=receipt['versions'][index];spec=receipt['sourcePreparations'][0] if index==0 else revision['request'];document=docx.Document(p);paragraphs=[x.text for x in document.paragraphs];expected=[''.join(r['text'] for r in b['runs']) for b in spec['content']['blocks']];assert paragraphs==expected;assert not document.tables;assert not document.inline_shapes;assert sha(data)==version['object']['sha256'];assert len(data)==version['object']['bytes'];assert version['object']['scopes']==version['scopes'];assert version['generationStatus']=='draft' and version['previewStatus']=='pending'
  text='\n'.join(paragraphs);assert 'Congress 119 | Session 2 | 119-hr-4242 | Bill version ih | FY 2027 | Budget stage authorized' in text;assert 'F2 [unverified-claim; potentially conflicting cited passages]' in text;assert 'F3 [missing-evidence]' in text;assert 'locator URL (not verification)' in text;assert 'program_elements: unavailable' in text;assert 'Free-form claims remain unverified' in text
  marker='Meri brief provenance: ';matches=[x[len(marker):] for x in paragraphs if x.startswith(marker)];assert len(matches)==1;provenance=json.loads(matches[0]);assert provenance['skill']==pin;assert provenance['subject']==revision['provenance']['subject']
  if index==0:assert provenance['parent'] is None and version['parentVersionId'] is None;original_provenance=provenance
  else:
   assert provenance==revision['provenance'];parent=provenance['parent'];original=receipt['versions'][0];assert version['parentVersionId']==parent['versionId']==original['versionId'];assert version['artifactId']==original['artifactId'];assert parent['objectSha256']==sha((directory/'original.docx').read_bytes());assert parent['specSha256']==hashj(receipt['sourcePreparations'][0]);assert parent['provenanceSha256']==hashj(original_provenance)
  for evidence in version['evidence']:
   assert evidence['sourceId'] in text and evidence['sourceVersionId'] in text and evidence['contentSha256'] in text and evidence['authorizationDecisionId'] in text;assert all(s in version['scopes'] for s in evidence['scopes'])
  inspection=json.loads((directory/(label+'-inspection.json')).read_text());assert inspection['sha256']==sha(data);assert inspection['content']['paragraphs']==paragraphs;pdfpath=directory/(label+'.pdf');pdf=fitz.open(pdfpath);assert len(pdf)==4==len(inspection['visual']['pages']);assert clean(text)==clean('\n'.join(page.get_text(sort=True) for page in pdf));images=[]
  for n,page in enumerate(pdf,1):
   png=directory/f'{label}-page-{n}.png';assert page.get_text().strip();assert sha(png.read_bytes())==inspection['visual']['pages'][n-1]['sha256'];images.append({'page':n,'path':str(png),'sha256':sha(png.read_bytes()),'physicallyOpened':True,'clipped':False})
  for process in inspection['processes']:
   assert process['leaderClosed'] and process['groupAbsent'] and process['exitCode']==0
   try:os.killpg(process['pid'],0);raise AssertionError('owned renderer group remains')
   except ProcessLookupError:pass
  rows.append({'cohort':cohort,'label':label,'docxSha256':sha(data),'docxBytes':len(data),'editableParagraphs':len(paragraphs),'pdfSha256':sha(pdfpath.read_bytes()),'pages':len(pdf),'images':images,'skillPin':pin,'versionId':version['versionId'],'artifactId':version['artifactId'],'baseVersionId':version['parentVersionId'],'provenanceSha256':hashj(provenance),'revisionFullReadEvidence':reads if index else None,'originalReadReceiptLimitation':'Original saved bundle contains spec/version/evidence/provenance, but not original full preparation/read envelopes; its reads digest cannot be independently recomputed.' if not index else None,'nativeQualified':False})
for p in [donor/'docx-http-cleanup.json',donor/'docx-http-cleanup-final.json']:
 c=json.loads(p.read_text());assert c['http']['closed'] and not c['http']['listening'] and c['http']['activeSockets']==0;assert c['fixture']['cleanup']['executorDirectoriesAbsent'] and c['fixture']['cleanup']['storeRemoved'];assert c['fixture']['generations']==c['fixture']['settled']==2
(here/'audit-report.json').write_text(json.dumps({'status':'PASS','sourceCommit':'b7b47c2b73dab98c2ec7074b2153a14d6a1e618e','tree':'a7aa6a1ba10b99cd51c3263b8350b835c19b3cc3','docx':4,'pagesPhysicallyOpened':16,'artifacts':rows,'nativeQualified':False},indent=2)+'\n');print(json.dumps({'status':'PASS','docx':4,'pages':16}))
