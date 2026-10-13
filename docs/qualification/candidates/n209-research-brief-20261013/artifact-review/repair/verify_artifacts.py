"""Independent offline reads of exact synthetic DOCX/PDF/PNG and source provenance.

Usage: python3 verify_artifacts.py [artifact-directory]. Exit 0 pass; 1 failed assertion.
Writes no artifacts; one bounded run is captured by the review owner.
"""
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import subprocess
import sys
from zipfile import ZipFile

import docx
from docx import Document
import fitz

ROOT = Path(__file__).resolve().parents[6]
EVIDENCE = Path(__file__).resolve().parent
ARTIFACTS = Path(sys.argv[1]).resolve() if len(sys.argv) == 2 else EVIDENCE.parent.parent / "integration/repaired-final-artifacts"
assert len(sys.argv) in (1, 2)


def digest(value):
    return sha256(value).hexdigest()


def canonical(value):
    # This fixture uses only exact integral JSON numbers, no floats or lone surrogates.
    if isinstance(value, float):
        raise AssertionError("Independent canonical JSON supports this integral fixture only")
    if isinstance(value, list):
        for item in value:
            canonical(item)
    if isinstance(value, dict):
        for item in value.values():
            canonical(item)
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()


def hash_json(value):
    return digest(canonical(value))


def compact(value):
    return re.sub(r"\s+", "", value)


def load(path):
    return json.loads(path.read_text())


pin = "57c8f7b99043938666b8f29f1e61e81c96c81d39"
head = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
tree = subprocess.check_output(["git", "rev-parse", f"{pin}^{{tree}}"], cwd=ROOT, text=True).strip()
assert tree == "b2729e85e4bbf3647129128edabb455d6847f96b"
provenance_file = ARTIFACTS / "provenance.json"
metadata = load(provenance_file)
versions = metadata["versions"]
preparations = metadata["preparations"]
assert len(versions) == len(preparations) == 2
manifest_path = ROOT / "apps/control/document-application/research-brief/manifest.json"
prose_path = manifest_path.with_name("SKILL.md")
catalog_path = ROOT / "packages/contracts/retrieval-0.2/catalog.json"
manifest = load(manifest_path)
catalog = load(catalog_path)
prose = prose_path.read_bytes()
skill_pin = {"skillId": manifest["skillId"], "skillVersion": manifest["skillVersion"],
             "manifestSha256": hash_json(manifest), "proseSha256": digest(prose),
             "catalogSha256": hash_json(catalog)}
assert manifest["prose"]["sha256"] == digest(prose)
assert skill_pin["skillId"] == "quick_skill.bill_watch" and skill_pin["skillVersion"] == "1.0.0"
subject = {"congress": 119, "session": 2, "billId": "119-hr-4242", "billVersion": "ih", "fiscalYear": 2027, "budgetStage": "authorized"}
summaries = []
source_receipts = []
section_properties = []
styles = []
texts = []
all_processes = []
for index, label in enumerate(["original", "revision"]):
    version, prepared = versions[index], preparations[index]
    request, source_reads, facts, provenance = [prepared[key] for key in ["request", "reads", "facts", "provenance"]]
    assert provenance["skill"] == skill_pin and provenance["subject"] == subject
    assert hash_json(source_reads) == provenance["readsSha256"]
    assert request["operation"] == "create_word" and request["format"] == "docx"
    expected_paragraphs = ["".join(run["text"] for run in block["runs"]) for block in request["content"]["blocks"]]
    assert len(expected_paragraphs) == len(request["content"]["blocks"])
    proposed = [text for text in expected_paragraphs if text.startswith(("Proposed:", "Revised proposal:"))]
    assert len(proposed) == 1
    reconstructed_input = {"skill": skill_pin, "title": request["title"], "subject": subject, "asOf": provenance["asOf"],
                           "facts": [{key: fact[key] for key in ["id", "kind", "text", "citations"]} for fact in facts],
                           "proposedLanguage": proposed}
    assert hash_json(reconstructed_input) == provenance["inputSha256"]
    source_evidence = []
    for read in source_reads:
        envelope = read["envelope"]
        assert envelope["method"] == "source-read"
        assert envelope["binding"] == version["binding"]
        assert envelope["subject"]["scopes"] == version["scopes"]
        assert envelope["subject"]["tenantId"] == version["binding"]["tenantId"]
        assert envelope["subject"]["actorId"] == version["binding"]["actorId"]
        assert envelope["subject"]["entityIds"] == [version["scopes"][0]["resourceId"]]
        selector = {"subject": subject, "asOf": provenance["asOf"]}
        assert hash_json(selector) == envelope["argumentsSha256"]
        for name, value in subject.items():
            name = re.sub(r"[A-Z]", lambda m: "_" + m.group().lower(), name)
            matches = [entry for entry in envelope["appliedFilters"] if entry["name"] == name]
            assert matches == [{"name": name, "valueSha256": hash_json(value)}]
        records = (envelope.get("data") or {}).get("records", [])
        assert len(records) == len(read["evidence"])
        assert len(envelope["sources"]) == 1
        source = envelope["sources"][0]
        source_id = source["sourceId"]
        source_receipts.append({"sourceId": source_id, "selectorSha256": hash_json(selector), "resultSha256": hash_json(read)})
        if source_id == "program_elements":
            assert not records and not read["evidence"] and envelope["status"] == "unavailable"
            assert source_id in envelope["coverage"]["sourcesUnavailable"]
            assert envelope["negativeAssertion"]["kind"] == "none"
        else:
            assert envelope["status"] == "complete" and len(records) == 1
        for record in records:
            assert record["content"]["kind"] == "full-read"
            evidence = next(ev for ev in read["evidence"] if ev["citationId"] == record["citationId"])
            assert evidence["contentSha256"] == record["contentSha256"] == digest(record["content"]["text"].encode())
            assert evidence["recordId"] == record["recordId"] and evidence["sourceVersionId"] == record["recordVersionId"]
            assert evidence["scopes"] == record["scopes"] == version["scopes"]
            assert evidence["authorizationDecisionId"] == envelope["subject"]["authorizationDecisionId"]
            assert evidence["observedAt"] == source["retrievedAt"]
            source_evidence.append(evidence)
    assert source_evidence == version["evidence"]
    assert [fact["status"] for fact in facts] == ["bound-excerpt", "unverified-claim", "missing-evidence"]
    assert facts[1]["kind"] == "claim" and facts[1]["conflict"] is True
    # Compare actual read text/ranges and citations, not the author's invariant flags.
    for fact in facts:
        for link in fact["links"]:
            if link["status"] != "read-and-bound":
                assert link["evidence"] is None and link["passage"] is None
                continue
            citation, evidence = link["citation"], link["evidence"]
            read = next(read for read in source_reads if read["envelope"]["sources"][0]["sourceId"] == citation["sourceId"])
            record = next(record for record in read["envelope"]["data"]["records"] if record["recordId"] == citation["recordId"])
            # UTF-16 unit selection is decoded independently, including its exact byte hash.
            raw_units = record["content"]["text"].encode("utf-16-le")
            passage = raw_units[citation["start"] * 2:citation["end"] * 2].decode("utf-16-le")
            assert passage == link["passage"] and digest(passage.encode()) == link["passageSha256"]
            assert citation["sourceVersionId"] == record["recordVersionId"]
            assert evidence in read["evidence"]
            assert f"Cited passage (UTF-16 characters {citation['start']}-{citation['end']}): {passage}" in expected_paragraphs
            identity = f"Source {evidence['sourceId']}; record {evidence['recordId']}; version {evidence['sourceVersionId']}; citation {evidence['citationId']}; authorization decision {evidence['authorizationDecisionId']}."
            assert identity in expected_paragraphs
            scope_hash_block = next(text for text in expected_paragraphs if text.startswith("Full text SHA256: " + evidence["contentSha256"]))
            assert f"Passage SHA256: {link['passageSha256']}" in scope_hash_block
            assert json.loads(scope_hash_block.split("Scope: ")[1].split("\n")[0]) == evidence["scopes"]
    raw = (ARTIFACTS / f"{label}.docx").read_bytes()
    assert digest(raw) == version["object"]["sha256"] and len(raw) == version["object"]["bytes"]
    document = Document(BytesIO(raw))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs == expected_paragraphs
    assert not document.tables
    assert document.paragraphs[0].style.name == "Title"
    assert [paragraph.style.name for paragraph in document.paragraphs if paragraph.text in ["Facts and evidence", "Proposed language", "Coverage and verification appendix"]] == ["Heading 1"] * 3
    assert paragraphs == load(ARTIFACTS / f"{label}-content.json")["paragraphs"]
    with ZipFile(BytesIO(raw)) as package:
        assert "word/document.xml" in package.namelist()
        xml = package.read("word/document.xml")
        assert b"<w:t" in xml and b"<w:drawing" not in xml
        assert b"w:documentProtection" not in package.read("word/settings.xml")
        styles.append(digest(package.read("word/styles.xml")))
    section_properties.append(document.sections[0]._sectPr.xml)
    edit_marker = "Independent editable body proof; retained original unchanged."
    document.paragraphs[0].text = edit_marker
    edited = BytesIO()
    document.save(edited)
    assert Document(BytesIO(edited.getvalue())).paragraphs[0].text == edit_marker
    assert digest((ARTIFACTS / f"{label}.docx").read_bytes()) == digest(raw)
    text = "\n".join(paragraphs)
    texts.append(text)
    assert "locator URL (not verification)" in text
    assert "F2 [unverified-claim; potentially conflicting cited passages]" in text
    assert "F3 [missing-evidence]" in text
    assert "Draft proposals for human review; not enacted language or verified facts." in text
    parsed_provenance = json.loads(next(p for p in paragraphs if p.startswith("Meri brief provenance: ")).split("Meri brief provenance: ")[1])
    assert parsed_provenance == provenance
    pdf_raw = (ARTIFACTS / f"{label}-preview.pdf").read_bytes()
    inspection = load(ARTIFACTS / f"{label}-inspection.json")
    assert inspection["nativeQualified"] is False and inspection["localCopiesOnly"] is True
    assert inspection["sourceSha256"] == digest(raw)
    assert inspection["transformReceipt"]["inputSha256"] == digest(raw)
    assert inspection["transformReceipt"]["outputSha256"] == digest(pdf_raw)
    assert inspection["visual"]["sha256"] == digest(pdf_raw)
    pdf_pages = []
    with fitz.open(stream=pdf_raw, filetype="pdf") as pdf:
        assert pdf.page_count == len(inspection["visual"]["pages"]) == 4
        assert not pdf.is_encrypted
        pdf_text = "\n".join(page.get_text() for page in pdf)
        assert compact(pdf_text) == compact(text), f"All DOCX body characters must survive PDF rendering ({label})"
        for page_index, page in enumerate(pdf):
            png_path = ARTIFACTS / f"{label}-preview-{page_index + 1}.png"
            png = png_path.read_bytes()
            image = fitz.Pixmap(png)
            recorded = inspection["visual"]["pages"][page_index]
            assert page.get_text() == recorded["text"]
            assert digest(png) == recorded["sha256"] and len(png) == recorded["bytes"]
            assert image.width == recorded["width"] == 765 and image.height == recorded["height"] == 990
            assert page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False).tobytes("png") == png
            # Independently inspect every text span's geometry within the actual page.
            spans = [span for block in page.get_text("dict")["blocks"] if block["type"] == 0 for line in block["lines"] for span in line["spans"]]
            assert spans
            assert all(page.rect.contains(fitz.Rect(span["bbox"])) for span in spans)
            pdf_pages.append({"page": page_index + 1, "pngSha256": digest(png), "pngBytes": len(png), "textSha256": digest(page.get_text().encode()), "spanCount": len(spans), "allTextSpansWithinPage": True})
    for process in inspection["processes"]:
        assert process["exitCode"] == 0 and process["leaderClosed"] is True and process["groupAbsent"] is True
        assert process["stdoutEnded"] is True and process["stderrEnded"] is True and not process["streamErrors"]
        assert process["sandboxQualified"] is False
        all_processes.append(process["pid"])
    summaries.append({"label": label, "versionId": version["versionId"], "docxSha256": digest(raw), "docxBytes": len(raw),
                      "pdfSha256": digest(pdf_raw), "pdfBytes": len(pdf_raw), "pageCount": len(pdf_pages), "pages": pdf_pages,
                      "paragraphs": len(paragraphs), "editableTextSaveAndReopen": True, "bodyMatchesPreparedRequest": True,
                      "bodyMatchesIndependentPdfTextIgnoringWhitespace": True, "stylesSha256": styles[-1], "specSha256": hash_json(request)})
assert source_receipts == metadata["sourceReads"]
assert styles[0] == styles[1] and section_properties[0] == section_properties[1]
original_provenance, revision_provenance = [prepared["provenance"] for prepared in preparations]
assert original_provenance["parent"] is None and versions[0]["parentVersionId"] is None
parent = revision_provenance["parent"]
assert versions[0]["artifactId"] == versions[1]["artifactId"]
assert versions[0]["versionId"] != versions[1]["versionId"]
assert versions[1]["parentVersionId"] == parent["versionId"] == versions[0]["versionId"]
assert parent["objectSha256"] == summaries[0]["docxSha256"]
assert parent["specSha256"] == hash_json(preparations[0]["request"])
assert parent["provenanceSha256"] == hash_json(original_provenance)
assert original_provenance["asOf"] <= revision_provenance["asOf"]
assert "Proposed: Request confirmation" in texts[0] and "Revised proposal:" not in texts[0]
assert "Revised proposal: Require independent confirmation" in texts[1] and "Proposed: Request confirmation" not in texts[1]
result = {"status": "PASS", "reviewedCommit": pin, "reviewedTree": tree, "observedHead": head,
          "startedAt": datetime.now(timezone.utc).isoformat(), "python": sys.executable, "pythonDocx": docx.__version__, "pyMuPdf": fitz.VersionBind,
          "artifactDirectory": str(ARTIFACTS), "provenanceFileSha256": digest(provenance_file.read_bytes()), "skillPins": skill_pin,
          "sourceReadBindings": source_receipts, "versions": summaries, "parent": parent, "sourceReopenProcessIds": all_processes,
          "limitations": ["Synthetic sources/authorization/identity/accounting only", "Local editable/read/render evidence only", "Hashes prove correspondence, not source authenticity or corpus completeness", "No native runtime/DB/provider/SDK checkpoint/browser qualification"],
          "nativeQualified": False, "productionActivated": False}
print(json.dumps(result, indent=2))
