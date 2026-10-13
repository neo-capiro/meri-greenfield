#!/usr/bin/env python3
"""Independently inspect synthetic editable artifacts, every PDF page and exact provenance."""
import hashlib,json,subprocess,sys,zipfile
from pathlib import Path
import docx,fitz

root=Path(__file__).resolve().parents[5]
directory=Path(sys.argv[1]).resolve()
output=Path(sys.argv[2]).resolve()
assert output.parent==Path(__file__).resolve().parent
sha=lambda value:hashlib.sha256(value).hexdigest()
canonical=lambda value:json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()
hashdata=lambda value:sha(canonical(value))
normal=lambda value:''.join(value.split())
utf16=lambda value:len(value.encode('utf-16-le'))//2
source=json.loads((directory/'provenance.json').read_text())
versions=source['versions'];preparations=source['preparations']
prose=(root/'packages/skills/originals/mail_history/1.0.0/SKILL.md').read_bytes()
manifest=json.loads((root/'packages/skills/originals/mail_history/1.0.0/manifest.json').read_text())
catalog=json.loads((root/'packages/contracts/retrieval-0.2/catalog.json').read_text())
rows=[]
for index,label in enumerate(['original','revision']):
 p=preparations[index];version=versions[index];raw=(directory/(label+'.docx')).read_bytes()
 assert sha(raw)==version['object']['sha256'] and len(raw)==version['object']['bytes']
 document=docx.Document(directory/(label+'.docx'))
 actual=[paragraph.text for paragraph in document.paragraphs]
 expected=[''.join(run['text'] for run in block['runs']) for block in p['request']['content']['blocks']]
 assert actual==expected, 'Editable DOCX paragraphs differ from retained typed request'
 assert not document.inline_shapes and not document.tables
 with zipfile.ZipFile(directory/(label+'.docx')) as archive:
  xml=archive.read('word/document.xml').decode()
  assert '<w:t' in xml and '<w:document' in xml and '<w:drawing' not in xml
 marker='Meri Mail History provenance: '
 embedded=[json.loads(paragraph[len(marker):]) for paragraph in actual if paragraph.startswith(marker)]
 assert embedded==[p['provenance']]
 pin=embedded[0]['skill'];assert pin['skillId']=='quick_skill.mail_history' and pin['skillVersion']=='1.0.0'
 assert pin['proseSha256']==sha(prose)==manifest['prose']['sha256']
 assert pin['manifestSha256']==hashdata(manifest) and pin['catalogSha256']==hashdata(catalog)
 assert embedded[0]['readsSha256']==hashdata(p['reads'])
 assert embedded[0]['selection']['subject']==p['reads'][0]['result']['subject']
 assert embedded[0]['selection']['subject']['scopes']==version['scopes']==p['context']['scopes']
 assert version['evidence']==p['context']['evidence']
 assert sorted({c['sourceId'] for c in p['coverage']})==['entity_candidates','mail','tenant_contacts']
 assert all(c['verified'] is False and c['enumeration']=='unknown' for c in p['coverage'])
 assert all('[unverified interpretation]' in value for value in actual if value.startswith('[unverified interpretation]'))
 decisive=[]
 for body in p['bodies']:
  if not body['evidence']:continue
  candidate=body['candidate'];sections=[read['result'] for read in p['reads'] if read['operation']['id']=='mail.read_thread' and read.get('result',{}).get('messages',[{}])[0].get('messageId')==candidate['messageId']]
  assert len(sections)==len(body['sections'])>1
  stitched='';offset=0;identities=[]
  for page in sections:
   row=page['messages'][0];section=row['section']
   assert page['binding']==p['context']['binding'] and page['window']==embedded[0]['selection']['window']
   assert page['subject']['scopes']==p['context']['scopes'] and row['scopes']==candidate['scopes']
   assert page['sources'][0]['sourceId']=='mail' and page['sources'][0]['authority']=='authoritative'
   assert row['readHandle'] in page['continuation']['requiredReads']
   assert section['unit']=='utf16-code-unit' and section['start']==offset
   assert utf16(section['text'])==section['endExclusive']-section['start']
   offset=section['endExclusive'];stitched+=section['text']
   identities.append({key:row[key] for key in ['sourceId','mailboxId','threadId','messageId','recordVersionId','citationId','contentSha256','scopes','bodyCompleteness']})
  assert all(identity==identities[0] for identity in identities)
  assert stitched==body['text'] and offset==utf16(stitched)==sections[-1]['messages'][0]['section']['totalCodeUnits']
  assert sections[-1]['messages'][0]['section']['hasMore'] is False
  evidence=body['evidence'];assert sha(stitched.encode())==evidence['contentSha256']==identities[0]['contentSha256']
  assert evidence['sourceVersionId']==identities[0]['recordVersionId'] and evidence['authorizationDecisionId']==sections[-1]['subject']['authorizationDecisionId']
  for quote in p['quotes']:
   if quote['status']!='bound-source-excerpt' or quote['messageId']!=candidate['messageId']:continue
   excerpt=stitched.encode('utf-16-le')[2*quote['start']:2*quote['end']].decode('utf-16-le')
   assert excerpt==quote['text']==quote['passage'] and sha(excerpt.encode())==quote['passageSha256']
   assert quote['start']>2000 and quote['evidence']==evidence
   decisive.append({'messageId':candidate['messageId'],'sourceVersionId':evidence['sourceVersionId'],'bodySha256':evidence['contentSha256'],'passageSha256':quote['passageSha256'],'utf16Range':[quote['start'],quote['end']],'sections':len(sections),'authorizationDecisionId':evidence['authorizationDecisionId']})
 assert decisive
 pdfbytes=(directory/(label+'-preview.pdf')).read_bytes();pdf=fitz.open(stream=pdfbytes,filetype='pdf')
 receipt=json.loads((directory/(label+'-inspection.json')).read_text())
 assert len(receipt['visual']['pages'])==len(pdf)>0 and receipt['sourceSha256']==sha(raw) and receipt['nativeQualified'] is False
 pdftext='\n'.join(page.get_text() for page in pdf)
 assert normal('\n'.join(actual))==normal(pdftext), 'Rendered PDF omitted or altered editable text'
 pages=[]
 for number,page in enumerate(pdf,1):
  png=directory/f'{label}-preview-{number}.png';expected_page=receipt['visual']['pages'][number-1]
  assert expected_page['page']==number and sha(png.read_bytes())==expected_page['sha256']
  assert normal(page.get_text())==normal(expected_page['text'])
  bounds=page.rect
  for word in page.get_text('words'):
   assert word[0]>=bounds.x0-1 and word[1]>=bounds.y0-1 and word[2]<=bounds.x1+1 and word[3]<=bounds.y1+1, 'Rendered text outside page bounds'
  pages.append({'page':number,'png':str(png),'sha256':sha(png.read_bytes()),'textCharacters':len(page.get_text())})
 assert all(process['exitCode']==0 and process['leaderClosed'] and process['groupAbsent'] for process in receipt['processes'])
 rows.append({'label':label,'versionId':version['versionId'],'artifactId':version['artifactId'],'parentVersionId':version['parentVersionId'],'docxSha256':sha(raw),'docxBytes':len(raw),'editableParagraphs':len(actual),'pdfSha256':sha(pdfbytes),'pages':pages,'decisivePassages':decisive,'skill':pin,'provenanceSha256':hashdata(p['provenance']),'nativeQualified':False})
assert versions[0]['parentVersionId'] is None
assert versions[1]['parentVersionId']==versions[0]['versionId'] and versions[1]['artifactId']==versions[0]['artifactId']
parent=preparations[1]['provenance']['parent'];assert parent['versionId']==versions[0]['versionId'] and parent['objectSha256']==versions[0]['object']['sha256']
assert parent['specSha256']==hashdata(preparations[0]['request']) and parent['provenanceSha256']==hashdata(preparations[0]['provenance'])
receipt={'status':'PASS','sourceHead':subprocess.check_output(['git','rev-parse','HEAD'],cwd=root,text=True).strip(),'artifactDirectory':str(directory),'docxLibrary':docx.__version__,'pdfLibrary':fitz.__version__,'versions':rows,'everyPageRendered':True,'visualInspectionRequiredSeparately':True,'nativeQualified':False}
output.write_text(json.dumps(receipt,indent=2)+'\n')
print(json.dumps(receipt,indent=2))
