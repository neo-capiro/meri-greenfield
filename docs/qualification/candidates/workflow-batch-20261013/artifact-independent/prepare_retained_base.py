#!/usr/bin/env python3
"""Reconstruct exact read-only base metadata from the retained real synthetic DOCX."""
import hashlib,json,zipfile
from pathlib import Path
from docx import Document
root=Path(__file__).resolve().parents[5]
directory=root/'apps/control/document-application/mail-history/local/receipts-v2'
saved=json.loads((directory/'provenance.json').read_text());p=saved['preparations'][0];version=saved['versions'][0]
raw=(directory/'original.docx').read_bytes();assert hashlib.sha256(raw).hexdigest()==version['object']['sha256'] and len(raw)==version['object']['bytes']
doc=Document(directory/'original.docx');assert [paragraph.text for paragraph in doc.paragraphs]==[''.join(r['text'] for r in b['runs']) for b in p['request']['content']['blocks']]
digest=lambda value:hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
with zipfile.ZipFile(directory/'original.docx') as archive:members=archive.namelist()
inspection={'sha256':version['object']['sha256'],'bytes':len(raw),'format':'docx','nativeOpened':False,'rendered':False,'recalculated':False,'members':members,'details':{'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'pictures':len(doc.inline_shapes)}}
snapshot={'version':version,'spec':p['request'],'specSha256':digest(p['request']),'requestSha256':digest({'request':p['request'],'context':p['context'],'baseVersionId':None}),'inspection':inspection}
(Path(__file__).resolve().parent/'retained-base.json').write_text(json.dumps({'snapshot':snapshot,'preparation':p,'realOriginalSha256':version['object']['sha256']},indent=2)+'\n')
print(json.dumps({'status':'PASS','realOriginalSha256':version['object']['sha256'],'paragraphs':len(doc.paragraphs),'members':len(members)}))
