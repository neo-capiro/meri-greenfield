#!/usr/bin/env python3
"""Read-only source/plan/workbook audit; writes only this lane's new JSON evidence."""
import copy
import hashlib
import json
import subprocess
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[5]
OUT = Path(__file__).resolve().parent
PIN = '18adda2c5461b8d346277135aafd3972b4768802'
TREE = 'eaec635b6b1c6cdc353ad48a5fda1f7a3dc10b21'
LEDGER = 'docs/qualification/candidates/vertical-source-20261012/integration/requirements-continuation.json'
PLAN = 'docs/plans/MERI-DEPLOYMENT.md'
BOOK = 'docs/plans/revisions/20261010-R2/Meri_Capability_Inventory-20261010-R2.xlsx'
MANIFEST = 'docs/plans/revisions/20261010-R2/manifest.json'
CATALOG = 'packages/contracts/retrieval-0.2/catalog.json'
SKILLS = 'docs/coordination/r2/skills.json'
BENCH = 'evals/retrieval-r2/manifest.json'
I = 'implemented-but-unqualified'
M = 'missing-implementation'
E = 'external-runtime'
D = 'genuinely-deferred'

# Each anchor is independently located in the actual frozen file, not an earlier
# report's generic source-test total. Group references remain source-only.
GROUPS = {
 'catalog': [
  (CATALOG, '"originals"'),
  ('packages/contracts/retrieval-0.2/validate.mjs', 'export const {assertCatalog}'),
  ('packages/contracts/retrieval-0.2/operation-contracts.mjs', 'export'),
  ('apps/gateway/catalog/runtime-catalog.mjs', 'export function createRuntimeCatalog'),
  ('apps/gateway/index.mjs', 'createSourceDisabledGateway')],
 'scope': [
  ('packages/contracts/retrieval-0.2/common.schema.json', '"SubjectRequest"'),
  ('packages/contracts/retrieval-0.2/validate.mjs', 'export function assertSubject'),
  ('apps/gateway/authorization/current-tool-authority.mjs', 'export function createCurrentToolAuthority'),
  ('apps/gateway/budget/admitted-run-owner.mjs', 'export function'),
  ('apps/control/action-host/resolver.mjs', 'export function createHostResolver'),
  ('apps/control/action-host/session/owner.mjs', 'export function')],
 'window': [
  ('packages/contracts/retrieval-0.2/common.schema.json', '"WindowRequest"'),
  ('packages/contracts/retrieval-0.2/validate.mjs', 'export function assertWindowResolution'),
  ('apps/control/document-application/research-brief/input.mjs', 'export'),
  ('apps/control/document-application/research-brief/README.md', 'No latest-version resolution')],
 'read': [
  ('packages/contracts/retrieval-0.2/validate.mjs', 'export function assertEvidence'),
  ('apps/gateway/dispatch/read-result-envelope.mjs', 'export function assertDispatchResult'),
  ('apps/gateway/sources/research-portfolio-evidence/handler.mjs', 'export function createResearchPortfolioEvidenceHandler'),
  ('apps/gateway/sources/research-portfolio-evidence/http-adapter.mjs', 'export'),
  ('apps/gateway/sources/action-items/handler.mjs', 'export'),
  ('apps/control/document-application/research-brief/evidence.mjs', 'export')],
 'skill': [
  (SKILLS, '"records"'),
  ('packages/contracts/retrieval-0.2/skill.schema.json', '"SkillInvocation"'),
  ('packages/contracts/retrieval-0.2/validate.mjs', 'export function assertSkillManifest'),
  ('apps/control/document-application/research-brief/SKILL.md', 'Identity: quick_skill.bill_watch'),
  ('apps/control/document-application/research-brief/manifest.json', '"skillId"'),
  ('apps/control/document-application/research-brief/skill.mjs', 'export'),
  ('apps/control/document-application/research-brief/README.md', 'cannot pass')],
 'budget': [
  ('apps/gateway/budget/tool-budget-execution.mjs', 'export'),
  ('apps/gateway/budget/tool-budget-dispatch.mjs', 'export'),
  ('apps/gateway/budget/postgres-tool-budget.mjs', 'export'),
  ('apps/gateway/budget/README.md', 'Settlement records'),
  ('apps/control/tool-invocations.mjs', 'export')],
 'native': [
  ('apps/worker/qualification/README.md', 'PREPARATION ONLY'),
  ('apps/worker/qualification/sdk-options.mjs', 'export'),
  ('apps/worker/qualification/checkpoint/resume-admission.mjs', 'export'),
  ('packages/contracts/qualification-sdk/native-evidence.mjs', 'export'),
  ('apps/control/qualification-temporal/README.md', 'Temporal')],
 'evaluation': [
  (BENCH, '"realRequestsReceived"'),
  ('docs/RETRIEVAL-AUDIT.md', 'real baseline uncollected'),
  ('docs/ROUTING-EVALUATION.md', 'dataset empty'),
  ('evals/task-lookup/corpus.mjs', 'export'),
  ('evals/task-lookup/score.mjs', 'export'),
  ('evals/task-lookup/execution/README.md', 'original'),
  ('evals/action-followthrough/local-baseline/README.md', 'not human-authored-case')],
 'replay': [
  ('apps/artifacts/recovery.mjs', 'export async function recoverArtifact'),
  ('apps/artifacts/postgres-service.mjs', 'ports.budget.recheck'),
  ('apps/artifacts/local-replay.mjs', 'artifact-replay'),
  ('apps/control/document-application/research-brief/compose.mjs', 'export'),
  ('docs/qualification/candidates/document-replay-20261013/artifact-review/final/report.md', '111/111 PASS'),
  ('docs/qualification/candidates/document-replay-20261013/source-review/recheck-report.md', '37/37 independent')],
 'optional': [
  ('docs/ROUTING-EVALUATION.md', 'Decisions/JEV remain disabled'),
  (BENCH, '"optionalDecisions"'),
  (BENCH, '"optionalJev"')]
}

# Primary class is a triage choice, not a claim that every subcriterion has that
# class. Every exact whole requirement remains pending; each bounded claim below
# names its observable source and its remaining boundary.
FINDINGS = {}
def record(rid, primary, *subs):
    FINDINGS[rid] = {'primaryClass': primary, 'subcriteria': [
        {'class': cls, 'criterion': criterion, 'finding': finding, 'evidenceGroups': refs}
        for cls, criterion, finding, refs in subs]}

record('R-01', I,
 (I, 'One typed catalog, domain contracts, preserved original mappings and availability guards', '139 original IDs are preserved: 96 selected all map to operation IDs/aliases and 43 deferred. 97 typed descriptors use canonical validators; runtime catalog checks schema/handler/descriptor pins and refuses unavailable entries.', ['catalog']),
 (M, 'Actual complete handler/availability/evaluation composition and unavailable-recipe agreement', 'Every production descriptor still has null handler/unimplemented/pending. Concrete source seams cover selected slices, not all 96 originals; original 14 recipe bodies cannot yet be linted against available handlers. Mapping completeness is not handler completeness.', ['catalog','skill','read']))
record('R-02', I,
 (I, 'Current actor/run/subject/grant boundaries at execution and result release', 'Canonical capability, current-run, business/source authorization and persisted admitted-run owners exist; session/control and artifact replay recheck current rights.', ['scope','replay']),
 (M, 'General same-name/pronoun/follow-up and authorized portfolio subject resolution', 'Typed SubjectRequest/ResolvedSubject and host-specific prepared scope do not implement a general trusted entity/pronoun resolver for every original skill.', ['scope','skill']),
 (E, 'Actual principals/policy/mail/memory/portfolio authorization and adversarial corpus', 'Authentic current identity/policy/source ports and original cross/within-tenant runtime suite are not established by local synthetic authority decisions.', ['scope','evaluation']))
record('R-03', M,
 (I, 'Typed calendar/duration/FY/version/stage contracts and supplied-boundary validation', 'WindowRequest separates calendar week/month/year, relative-duration, explicit and fiscal-year. Validators check supplied timezone/local-midnight calendar bounds and prevent basis/stage/revision substitution; research brief retains explicit legislative/FY/stage pins.', ['window']),
 (M, 'Trusted request-to-normalized-window and comparable domain basis resolution', 'No production resolver computes this/last week/month/year from a trusted clock and actor timezone, applies explicit Monday/Sunday policy, resolves fiscal-calendar bounds or authoritative revision/stage conflicts. assertWindowResolution checks supplied normalized bounds; it does not produce them.', ['window','scope']),
 (E, 'Authoritative actor preferences/fiscal calendars/version snapshots', 'Real policy/timezone and domain revision authority must be bound through existing trusted ports; synthetic supplied dates do not qualify the whole boundary suite.', ['window','evaluation']))
record('R-04', M,
 (I, 'Distinct candidate/exact/full-read shapes and bounded decisive passage compiler', 'Canonical typed domain schemas separate discovery and authoritative observations. Existing research/action source seams and brief compiler validate actual supplied decisive text, selectors, scope and authorization decisions.', ['read','window','replay']),
 (M, 'Authentic exact/search/full-read implementations for the entire selected source inventory', 'Disabled production facade plus contract bindings are not all selected source adapters. Complete multi-source history, hybrid candidate discovery, exact count/status/comparison and version/stage truth remain missing beyond the implemented slices.', ['catalog','read']),
 (E, 'Authorized source interfaces/snapshots and late-body/domain oracles', 'Authentic full text and authoritative producer versions/counts must come from qualified sources; synthetic passage/hash consistency is insufficient.', ['read','evaluation']))
record('R-05', I,
 (I, 'Coverage/freshness/version/continuation consistency and visible artifact appendix', 'Evidence envelopes and read/result validation preserve explicit coverage and lineage; brief includes passages, missing/conflicting evidence, verification appendix and immutable citations without granting free-form claims verified status from URLs.', ['read','replay']),
 (M, 'Complete source-specific cursor/snapshot/coverage implementation for all selected sources', 'Uniform DTO checks cannot prove truthful completeness, filters, freshness or stable cursors for source adapters not yet implemented/composed.', ['catalog','read']),
 (E, 'Actual producer authenticity and current citation/source authorization', 'Real producer/snapshot/current-source proof and complete enumeration remain unqualified; no hash/date syntax substitutes for source attestation.', ['read','evaluation']))
record('R-06', E,
 (I, 'Pinned SDK/proxy preparation, catalog discovery/availability and evidence contracts', 'SDK options/checkpoint/evidence source and default disabled catalog exist. They preserve Claude SDK as sole loop; source evidence is preparation or local qualification.', ['native','catalog']),
 (E, 'Actual model/runtime/proxy references, streaming, headers, compaction/resume and measured recall/cost/latency', 'No complete admitted supported native discovery-through-proxy run or adversarial required-capability recall benchmark is qualified by PR59.', ['native','evaluation']))
record('R-07', M,
 (I, 'Explicit 14 IDs/version/typed-input and manifest/pin continuity foundation', 'Canonical invocation/manifest validators exist; original bill_watch congressional recipe pins prose, manifest, catalog, input/read hashes and exact immutable parent.', ['skill','replay']),
 (M, 'Complete executable manifest selection/completion and lifecycle policy for all 14 skills', '13 original prose/manifest pairs are absent; the one brief recipe is a bounded output, not whole Bill Watch. General submit/steer/reconnect/resume/fork selection, active-version migration and mandatory-source completion need existing SDK/control/frontend composition.', ['skill','scope']),
 (E, 'Real UI/SDK/Temporal lifecycle and qualified operation availability', 'Structural pin tests and local artifact repeats do not qualify frontend selection or native SDK durable skill execution.', ['native','catalog','replay']))
record('R-08', M,
 (I, 'Preserved 14 original quick-action intents and bounded original congressional brief', 'The committed R2 skill inventory retains all 14 IDs including commitments_audit; product apps/packages have one original SKILL.md/manifest pair for bill_watch1.0.0.', ['skill']),
 (M, 'Individual original skills and promised evidence/output/completion rubrics', '13 original pairs are absent and complete Bill Watch remains pending. Commitments proof versus suggestion, unknown owner/date, evidence-backed completion, deduplication and mapped Monday coverage are requirements, not an implemented original skill procedure.', ['skill','read']),
 (E, 'Authorized required-source qualification for all promised workflows', 'Each skill must exercise genuine required-source coverage with current rights; no source widening or Quicklog UI is implied by filling original recipe gaps.', ['skill','catalog','evaluation']))
record('R-09', I,
 (I, 'Durable budgets, conservative unknown outcomes and bounded exact artifact recovery', 'Canonical admission/reservation/invocation/settlement owners enforce declared limits and preserve UNKNOWN. Exact replay retains original charge/version and current durable grant; no second physical generation is needed for verified completed results.', ['budget','replay']),
 (M, 'Requested-work complexity allocation and diagnosed source/model recovery policy', 'Declared accounting caps and direct typed lookup do not implement/measure general simple versus comprehensive task budgeting, token ceilings or purposeful fallback across the full source/skill inventory.', ['budget','skill','read']),
 (E, 'Actual SDK usage, settlement/reconciliation and paired task/token/latency observations', 'Fixed conservative charges and synthetic local receipts do not establish real measured model usage or end-to-end runtime enforcement.', ['budget','native','evaluation']))
record('R-10', M,
 (I, 'Redacted scoring/trace protocol and bounded correct-tool observation slice', 'Offline corpus/scorer and typed task-lookup observation producer retain pins, trace stage hashes, coverage and missing/ineligible denominators. Local six-case baseline is explicitly synthetic.', ['evaluation']),
 (M, 'Complete original/correct-tool/correct-evidence execution and route-to-final-claims diagnostics', 'Task-lookup executor supports correct-tool only; original and correct-evidence stay ineligible. General multi-domain/model/claim diagnostic execution is not implemented by that slice.', ['evaluation']),
 (E, '75 authenticated real requests, >=200 grouped cases, snapshots and actual timing/cost', 'Benchmark manifest has 0 real requests, empty case IDs, null source manifest/split/budgets/results. Original60-task and100-adversarial gates remain separate.', ['evaluation']))
record('R-11', D,
 (D, 'Optional bounded Decisions advisor and independent adoption', 'Disabled pending qualified SDK baseline, actual supported API and paired held-out calibration/fallback/latency/adoption. Optional work is not mandatory core completion or permanently excluded.', ['optional','evaluation']))
record('R-12', D,
 (D, 'Optional calibrated task-specific JEV evidence evaluation', 'Disabled pending frozen human-labelled relevance/support/client-fit calibration, abstention/safeguards and independent adoption after baseline. No unsupported router or second loop is required.', ['optional','evaluation']))
record('REL-01', E,
 (I, 'Source inventory/failure taxonomy/oracle and benchmark protocol', 'Audit and evaluation protocol are committed, with real-source absence disclosed rather than substituted by synthetic cases.', ['evaluation']),
 (E, '75 real seed traces, deployed settings/snapshots and measured baseline/budgets', 'Real seed and authenticated trace/source/config corpus remain unavailable; manifest still0. Archive transfer approval is not a real-request dataset or an approved provider benchmark.', ['evaluation']))
record('REL-02', I,
 (I, 'Frozen typed operations/scope/window/evidence/skill and original mapping foundation', 'One canonical139-original/97-operation catalog maps all96 selected aliases, retains43 deferred and validates unavailable activation. Public package/API bytes remain those accepted by prior owners.', ['catalog','scope','window','skill']),
 (M, 'Actual handler/schema/description/recipe agreement for full selected product', 'Production handler availability remains null/unimplemented/pending; 13 recipe pairs absent and full current available-handler evaluation mapping/completion are not supplied by schema publication alone.', ['catalog','skill']))
record('REL-03', M,
 (I, 'Exact/search/read contracts and independently written concrete source slices', 'Research-portfolio evidence, action items and related gateway projection/HTTP source boundaries exist, alongside typed domain contracts and full-read brief compilation.', ['read','catalog']),
 (M, 'All selected adapters, trusted windows and complete memory/client/portfolio/domain behavior', 'The source-disabled facade and pending97 descriptors do not compose all96 original capabilities. Trusted window normalization and missing source-family product implementations still need canonical source owners.', ['catalog','window','read']),
 (E, 'W03 qualified authentic current source checks and domain snapshots', 'Independent local source fixtures do not establish genuine source access/authorization/freshness/paging/oracle results.', ['read','evaluation']))
record('REL-04', M,
 (I, 'Typed selected skill identity foundation and one original brief output', 'All14 IDs remain in contracts/inventory; brief pins original1.0.0 manifest/prose and exact revision provenance using existing artifact lifecycle.', ['skill','replay']),
 (M, '14 complete original skills/Commitments and UI lifecycle/completion composition', '13 original recipe pairs plus the rest of Bill Watch, common manifest execution/required-source completion and requested window/portfolio/memory regressions remain product implementation work.', ['skill','window','scope']),
 (E, 'Frontend Playwright submit/steer/reconnect/resume/fork/native binding', 'This repo/source-only audit cannot claim independent frontend/native acceptance. Frontend owner and admitted runtime must prove actual selected identity and completion.', ['native','skill']))
record('REL-05', E,
 (I, 'Sole Claude SDK loop preparation and bounded source checkpoint/controls', 'SDK/proxy/runtime evidence source exists; budgets and exact source adapters retain canonical ownership rather than adding an orchestrator.', ['native','budget']),
 (M, 'Full lightweight/complexity/recovery integration across original skills and sources', 'Existing typed/local control slices do not supply missing common window/skill/source product behavior required for the later W09 checkpoint.', ['skill','window','read','budget']),
 (E, 'Runtime-only W02 proxy checkpoint then actual W03/W05/W09 integration', 'Native model/proxy/discovery/headers/stream/compaction/resume/adversarial recall and measured latency are still unqualified. Runtime-only checkpoint precedes source integration; no dependency circularity is introduced.', ['native','evaluation']))
record('REL-06', E,
 (I, 'Corpus freezing/scoring and limited observation/execution source', 'Existing evaluator protects grouping/gold/pins/denominators and correct-tool captures; synthetic baseline does not acquire real inputs.', ['evaluation']),
 (M, 'Complete oracle execution and all domain/skill/route diagnostic modes', 'Original and correct-evidence executions remain ineligible in task-lookup slice; broader200-case runtime harness and original skill/domain behaviors are not closed by deterministic scoring source.', ['evaluation','skill','read']),
 (E, 'Frozen >=200 actual cases, gold snapshots/baseline results and independent adjudication', 'Real seed0, no authenticated source snapshot or frozen split, no measured baseline/held-out results. REL01/03/04/05 and original gates remain prerequisites.', ['evaluation']))
record('REL-07', D,
 (D, 'Optional Decisions shadow after REL06 baseline', 'Remain disabled until qualified baseline and supported API; paired quality/coverage/argument/latency and independent adopt-or-disabled decision are pending.', ['optional','evaluation']))
record('REL-08', D,
 (D, 'Optional JEV shadow after REL06 baseline', 'Remain disabled until human-label/calibration/required-evidence safeguards and paired quality/latency evidence support independent adoption.', ['optional','evaluation']))
record('REL-09', E,
 (I, 'Recoverable reviewed immutable source candidates and bounded artifact/replay evidence', 'PR59 integrates source reviews including exact retained DOCX versions, current session/durable-grant reauthorization and unchanged public lifecycle. This is source/local acceptance only.', ['replay','catalog']),
 (M, 'Remaining required skill/window/source/oracle implementation', 'REL03/04/05/06 remaining product behaviors still prevent complete prestaging W09 evidence.', ['skill','window','read','evaluation']),
 (E, 'W09 integrated native/full-RG evidence, W10 immutable staging candidate, later authorized W11 canary', 'No deployment/canary or runtime gate is granted by PR59. Canary follows W10 and separate release authority; it is not a prerequisite for W09.', ['native','evaluation']))
record('RG-01', E,
 (I, 'Typed route/availability controls and score denominator protocol', 'Catalog refuses unavailable routes and offline scoring preserves missing/ineligible cases.', ['catalog','evaluation']),
 (M, 'Complete original request/model routing across all required source families and clarifications', 'The typed correct-tool lookup slice does not run original model routing; clarification observation remains ineligible and missing skills/sources limit complete route behavior.', ['evaluation','skill','read']),
 (E, '>=95% accepted held-out routes with all required families/expected clarifications', 'No frozen actual200-case suite or independent SDK baseline result; synthetic control totals are not this denominator.', ['evaluation','native']))
record('RG-02', M,
 (I, 'Deterministic supplied subject/window/FY/version checks', 'Canonical validators and explicit brief basis pins provide source controls, including timezone/DST consistency on supplied normalized bounds.', ['scope','window','replay']),
 (M, 'Trusted entity/window/fiscal/version resolution and full100% fixture behavior', 'General subject/pronoun and calendar/FY resolver plus all skill/domain conflict behavior are absent; validator checks are not resolver output tests.', ['scope','window','skill']),
 (E, 'Actual current subject/domain authority and exact full fixture suite', 'Real identity/source snapshots and authoritative revision/FY/stage proof remain unqualified.', ['scope','evaluation']))
record('RG-03', E,
 (I, 'Candidate/evidence DTOs and preservation of full gold/denominators', 'Source schemas and protocol retain search-versus-completeness distinction; evaluator does not shrink gold or grant empty-cohort success.', ['read','evaluation']),
 (M, 'Full discovery retrieval and atomic-unit Recall@10 execution/scoring across sources', 'Current action recommendation scorer is not a complete answer-bearing discovery Recall@10 benchmark; missing selected discovery adapters remain.', ['catalog','read','evaluation']),
 (E, 'Labelled frozen <=10-gold-unit Recall@10>=95%', 'No authenticated frozen candidate pools/gold/sources/results. >10-record histories must remain paging/completeness cases, exact counts RG04.', ['evaluation']))
record('RG-04', E,
 (I, 'Typed exact domain basis and supplied result consistency', 'Bill/PE/FY/stage/version contracts and read validators preserve domain distinctions; brief does not silently change requested budget stage/version.', ['window','read']),
 (M, 'All authoritative count/status/comparison adapters and trusted comparable basis resolver', 'Schemas alone do not perform every required exact query or resolve compatible subject/FY/stage/units/version for comparison.', ['catalog','window','read']),
 (E, '100% authorized snapshot/oracle deterministic domain suite', 'No full authenticated source snapshot oracle corpus and runtime results; exact Decimal/domain semantics must remain independently qualified.', ['evaluation','read']))
record('RG-05', I,
 (I, 'Incomplete coverage/candidate absence guards and visible brief verification limits', 'Canonical envelope validations and brief compiler preserve unavailable/partial/conflicting evidence; free-form URL-cited claims remain unverified.', ['read','replay']),
 (M, 'Final answer/completion enforcement for every original workflow', 'Missing14-skill completion/source composition means local envelope guards do not establish universal final absence-claim behavior.', ['skill','catalog']),
 (E, 'Zero false absence on full real partial/unavailable/stale/truncated/unpaged suite', 'No complete real-source/end-to-end model corpus or final-claim adjudication exists for this gate.', ['evaluation','native']))
record('RG-06', M,
 (I, '14 ID/version contract foundation and one exact original brief recipe', 'Original inventory/contract IDs are preserved; brief immutable manifest/prose/parent tests exist.', ['skill','replay']),
 (M, '100% all14 required evidence/completion/lifecycle fixtures', '13 pairs and full Bill Watch are absent; full client/portfolio/memory/Commitments submit/replay/steer/reconnect/resume/fork behavior is not implemented/qualified by this output slice.', ['skill','scope']),
 (E, 'Actual frontend and SDK/Temporal selected skill execution', 'Local artifact repeat is not UI skill submission or durable model replay acceptance.', ['native','evaluation']))
record('RG-07', I,
 (I, 'Current tenant/subject/session/resource and persisted grant controls', 'Gateway current authority, admitted run and artifact replay grant/resource/session checks fail closed; retained grant-only release counterexample and repair are visible.', ['scope','replay']),
 (E, 'Zero unauthorized disclosure/mutation original100 plus full source/context boundary qualification', 'Genuine principal/policy/source/storage and memory/mail/stored-context runtime suites remain required. Local synthetic decisions/test counts do not establish whole gate.', ['scope','native','evaluation']))
record('RG-08', E,
 (I, 'Native SDK/proxy evidence preparation and unavailable catalog fallback guard', 'Pinned preparation/source observations exist with deliberately disabled production capabilities.', ['native','catalog']),
 (E, 'Actual proxy tool references/deferred schemas/headers/streams/compaction/resume/adversarial recall', 'No complete admitted supported runtime/proxy and known-required-capability recall/fallback result is qualified; local PG client crash/prep does not supply it.', ['native','evaluation']))
record('RG-09', E,
 (I, 'Budget declarations and performance protocol', 'Fixed budget caps and source timing receipts exist; protocol distinguishes UIack/readiness/warm/cold/complete answer and preserves ordinary p95<=8s.', ['budget','evaluation']),
 (E, 'Frozen measured cohort/load, warm answer p95<=15s and overlapping original targets', 'No actual qualified model/native benchmark p50/p95 with full discovery/retry/advice/evidence latency. Optional<=500ms and<=10%regression applies only upon optional-service adoption.', ['evaluation','native','optional']))
record('RG-10', D,
 (D, 'Optional >=25% relative held-out route-error improvement with confidence/adjudication/latency', 'Unmeasured and disabled pending sufficiently powered paired baseline, no coverage/argument regression and RG09. It is not a core release blocker while disabled.', ['optional','evaluation']))
record('RG-11', D,
 (D, 'Optional labelled task-specific JEV improvement/calibrated abstention/no evidence loss', 'Unmeasured and disabled. Freeze separate relevance/support/client-fit/adoption thresholds before shadow trials; mandatory reads and authoritative facts remain protected.', ['optional','evaluation']))
record('RG-12', I,
 (I, 'Exact published document replay, revision lineage, current revocation and bounded decisive read', 'PR59 contains reviewed exact immutable original/revision and no-second-generation/charge recovery, full byte inspection, session/resource/durable-grant rechecks, visible decisive passages and preserved UNKNOWN for unsupported/unpublished gaps. Independent111 controls/23HTTP/8rendered pages are local synthetic source evidence.', ['replay','scope','read']),
 (M, 'All original skill/full-source completeness/reconnect/fork behavior', 'Document replay and congressional passage compilation are a bounded slice; missing skill/source/window product composition still prevents whole coverage/recovery contract completion.', ['skill','read','window']),
 (E, 'Actual SDK/Temporal/PostgreSQL/storage/identity/source lifecycle and original recovery gate', 'Local fs/HTTP/Office facts do not qualify native distributed durability, authentic source rights, rendering sandbox or all required recovery/reconnect actions.', ['native','evaluation','replay']))

def sha(data):
    return hashlib.sha256(data).hexdigest()

def git(*args):
    return subprocess.run(['git', *args], cwd=ROOT, capture_output=True, timeout=15)

def main():
    head = git('rev-parse','HEAD').stdout.decode().strip()
    tree = git('rev-parse','HEAD^{tree}').stdout.decode().strip()
    assert (head, tree) == (PIN,TREE), (head,tree)
    ledger = json.loads((ROOT/LEDGER).read_text())
    expected = {f'{p}-{n:02}' for p,m in [('R',12),('REL',9),('RG',12)] for n in range(1,m+1)}
    originals = {r['id']:r['governingRequirement'] for r in ledger['records'] if r['id'] in expected}
    assert set(FINDINGS) == set(originals) == expected and len(expected) == 33
    plan_lines = (ROOT/PLAN).read_text().splitlines()
    input_manifest = json.loads((ROOT/MANIFEST).read_text())
    assert sha((ROOT/PLAN).read_bytes()) == input_manifest['planSha256']
    assert sha((ROOT/BOOK).read_bytes()) == input_manifest['workbookSha256']
    old = ledger['governingSource']
    old_lookup = git('show', f"{old['commit']}:{old['path']}")
    # Keep failed original lookup as a failure, and use actual pinned R2 inputs.
    original_lookup = {'request':old, 'exitCode':old_lookup.returncode,
      'stdoutBytes':len(old_lookup.stdout),'stdoutSha256':sha(old_lookup.stdout),
      'stderr':old_lookup.stderr.decode(), 'boundary':'Original cbbd path is unavailable in this checkout; copied record hashes are preserved, not independently reconstructed from that unavailable object. All33 actual original plan lines and workbook rows are checked below.'}
    (OUT/'original-ledger-lookup.json').write_text(json.dumps(original_lookup,indent=2)+'\n')
    workbook = load_workbook(ROOT/BOOK,read_only=True,data_only=False)
    joins = {}
    for sheet in ['New Meri','Tool contracts','Deployment work','Qualification']:
      ws = workbook[sheet]
      headers = list(ws.iter_rows(min_row=7,max_row=7,values_only=True))[0]
      for row in ws.iter_rows():
        vals = [c.value for c in row]
        if vals and vals[0] in expected:
          rid=vals[0]
          joins.setdefault(rid,[]).append({'sheet':sheet,'row':row[0].row,
            'idCell':row[0].coordinate,'cells':{c.coordinate:c.value for c in row if c.value is not None},
            'fields':dict(zip(headers,vals)) if sheet!='New Meri' else None})
    workbook.close()
    assert set(joins)==expected
    records=[]
    for rid in sorted(expected):
      original=originals[rid]
      line=original['originalLocator']['line']
      assert original['originalLocator']['path']==PLAN
      assert plan_lines[line-1]==original['originalRequirement'],rid
      assert len(joins[rid])==(2 if rid.startswith('R-') else 1),rid
      assessment=copy.deepcopy(FINDINGS[rid])
      for sub in assessment['subcriteria']:
        assert sub['class'] in {I,M,E,D}
        assert sub['evidenceGroups'] and all(g in GROUPS for g in sub['evidenceGroups'])
      records.append({'id':rid,'governingRequirement':copy.deepcopy(original),
        'originalPlanCheck':{'path':PLAN,'line':line,'exactTextMatch':True},
        'originalWorkbookChecks':joins[rid],
        'wholeRequirementStatus':'pending','qualifiedRuntime':False,
        'assessment':assessment})
    evidence={}
    for group,files in GROUPS.items():
      evidence[group]=[]
      for path,anchor in files:
        data=(ROOT/path).read_bytes()
        lines=data.decode().splitlines()
        matches=[n for n,line in enumerate(lines,1) if anchor in line]
        assert matches,(path,anchor)
        evidence[group].append({'path':path,'sha256':sha(data),'bytes':len(data),
          'anchor':anchor,'line':matches[0]})
    catalog=json.loads((ROOT/CATALOG).read_text())
    selected={r['id'] for r in catalog['originals'] if r['disposition']=='selected'}
    mapped={alias for o in catalog['operations'] for alias in o.get('aliases',[])}|{o['operation']['id'] for o in catalog['operations']}
    assert len(catalog['originals'])==139 and len(selected)==96 and selected<=mapped
    assert all(o['handler'] is None and o['availability']=='unimplemented' and o['qualification']=='pending' for o in catalog['operations'])
    skill_rows=json.loads((ROOT/SKILLS).read_text())['records']
    skill_ids=[r['fields']['ID'] for r in skill_rows]
    assert len(skill_ids)==len(set(skill_ids))==14
    tracked=git('ls-files','apps','packages').stdout.decode().splitlines()
    skill_files=[p for p in tracked if Path(p).name=='SKILL.md']
    manifest_files=[p for p in tracked if Path(p).name=='manifest.json']
    assert skill_files==['apps/control/document-application/research-brief/SKILL.md']
    assert manifest_files==['apps/control/document-application/research-brief/manifest.json']
    brief=json.loads((ROOT/manifest_files[0]).read_text())
    assert brief['skillId']=='quick_skill.bill_watch'
    bench=json.loads((ROOT/BENCH).read_text())
    assert bench['realRequestsReceived']==0 and bench['caseIds']==[] and bench['results'] is None
    no_runtime=all(not r['qualifiedRuntime'] and r['wholeRequirementStatus']=='pending' for r in records)
    facts={'originalCatalog':{'originals':139,'selected':96,'deferred':43,'operations':len(catalog['operations']),
      'selectedMapped':len(selected & mapped),'selectedUnmapped':sorted(selected-mapped),
      'nullHandlerUnimplementedPending':len(catalog['operations'])},
      'originalSkills':{'selectedIds':skill_ids,'actualOriginalProseFiles':skill_files,
        'actualOriginalManifestFiles':manifest_files,'implementedBoundedRecipe':brief['skillId'],
        'absentOriginalPairs':[s for s in skill_ids if s!=brief['skillId']],
        'boundary':'One bounded congressional brief recipe is not complete Bill Watch or all14 qualification.'},
      'benchmark':bench,
      'sourceOnly':True,'nativeQualified':False,
      'currentCoordinationCorrections':[
       {'fact':'Neo approved private Meri archive/receipt transfer at10:16UTC; producer Library401 remains the blocker, not approval pending.',
        'authority':'Current parent task instruction; not independently executed connector access.','scope':'Exact authorized archive/receipt only; not other artifacts or75 real-request/model benchmarks.'},
       {'fact':'Separate native owner passed PR57 local install/prep/client-crash checks only; live worker/identity/activation remain blocked.',
        'authority':'Current parent task instruction; no independent native work in this audit.','scope':'Reported narrow native checks do not satisfy R06/REL05/RG08 or source authentication.'}]}
    counts={'records':len(records),'wholePending':33,'wholeRuntimeQualified':0,
      'primaryClassCounts':dict(Counter(r['assessment']['primaryClass'] for r in records)),
      'subcriterionClassCounts':dict(Counter(s['class'] for r in records for s in r['assessment']['subcriteria'])),
      'originalPlanExactJoins':len(records),'originalWorkbookRowJoins':sum(len(r['originalWorkbookChecks']) for r in records)}
    assert no_runtime and counts['wholeRuntimeQualified']==0
    audit={'schema':'meri-original-reliability-source-audit/1','source':{'commit':PIN,'tree':TREE},
      'boundary':'Independent bounded source audit; no new implementation/tests, private/native/provider/cloud action, whole-requirement promotion or runtime acceptance.',
      'classDefinitions':{
       I:'Concrete bounded source/contract behavior exists, with whole authentic integration/runtime/criteria still unqualified.',
       M:'Actual required product logic, full recipe/source composition or diagnostic execution is absent beyond documented/scaffolded inputs.',
       E:'An authentic external input, current authority/source snapshot, admitted native execution or measured benchmark/release action remains required.',
       D:'Only optional Decisions/JEV is deliberately disabled pending baseline and independent adoption; retained as optional scope.'},
      'originalInputs':{'ledgerPath':LEDGER,'ledgerSha256':sha((ROOT/LEDGER).read_bytes()),
        'plan':{'path':PLAN,'sha256':sha((ROOT/PLAN).read_bytes())},
        'workbook':{'path':BOOK,'sha256':sha((ROOT/BOOK).read_bytes())},
        'manifest':{'path':MANIFEST,'sha256':sha((ROOT/MANIFEST).read_bytes())},
        'unavailableOriginalLedgerLookup':'original-ledger-lookup.json',
        'joinBoundary':'Exact original plan lines verified against committed R2 hashes. Workbook row criteria/targets independently read; plan has more detailed qualifiers and current instructions supersede historical owner labels/statuses.'},
      'counts':counts,'facts':facts,'evidenceGroups':evidence,'records':records,
      'criticalPath':{
        'latestAuthorizedExecutionPriority':'Neo now prioritizes general conversation/event lifecycle with parallel original-skill/calendar workers, as supplied by the parent task. This audit does not override that instruction or implement product code.',
        'firstUnblockedProductIncrement':'Trusted typed WindowRequest-to-NormalizedWindow resolver through existing control/retrieval owners; no public package/API break or alternate orchestration.',
        'why':'R03 explicitly remains missing; original workbook asks resolve actor timezone, this week, FY/session/stage/exact revision. It is a shared prerequisite for R07/R08/REL03/REL04/RG02/RG04 and can be tested without private/provider/native operations.',
        'reuse':['packages/contracts/retrieval-0.2/common.schema.json','packages/contracts/retrieval-0.2/validate.mjs','apps/control/action-host/resolver.mjs','apps/gateway/authorization/current-tool-authority.mjs'],
        'acceptance':'Explicit trusted now/timezone/week policy; this/last week/month/year local-midnight UTC bounds through DST; true fixed relative durations; explicit intervals; fiscal-calendar authority plus FY/stage/revision consistency. Normalize supplied typed requests, retain known IDs/version/stage, refuse unresolved consequential authority, and exercise end-to-end existing consumer inputs. Validators remain validation, not a second resolver.',
        'secondProductIncrement':'Original versioned recipe/manifest/completion source for remaining13 intents plus complete Bill Watch, through existing Claude SDK/control/source/artifact owners. Prioritize required Commitments and a source-backed meeting brief; do not add Quicklog Commitment UI.',
        'parallelExternalCriticalPath':'Separate native owner supplies admitted identity/live worker/proxy/Temporal/storage/runtime; authorized archive producer fixes Library401 and transfers only approved receipt/archive. Corpus owner supplies authenticated75 real requests/snapshots, then>=200 frozen cases and model benchmark authorization.',
        'avoid':'No new qualification wrapper, generic registry, policy system, model loop, approval endpoint, new feature scope or synthetic substitution for original metrics.'}}
    (OUT/'audit.json').write_text(json.dumps(audit,indent=2,ensure_ascii=False)+'\n')
    validation={'decision':'ACCEPT_AUDIT_SOURCE_ONLY','source':audit['source'],'counts':counts,
      'checks':{'exactSourcePin':True,'originalInputHashes':True,'33OriginalPlanLines':True,
        '45OriginalWorkbookRows':True,'33UniqueAssessmentRecords':True,'evidencePathsHashesAnchors':True,
        '96SelectedOriginalMappings':True,'14OriginalSkillInventory':True,'zeroRealBenchmarkManifest':True,
        'noWholePromotion':True},'nativeQualified':False,
      'originalLedgerLookup':'Unavailable original cbbd path retained as failure; actual original plan/workbook criteria verified.',
      'auditSha256':sha((OUT/'audit.json').read_bytes())}
    (OUT/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
    print(json.dumps(validation,indent=2))

if __name__=='__main__':
    main()
