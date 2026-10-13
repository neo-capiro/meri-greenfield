"""Read-only local source and retained synthetic workbook audit."""
import os,json,pathlib,subprocess,hashlib
from openpyxl import load_workbook
root=pathlib.Path(__file__).resolve().parents[5];packet=pathlib.Path(__file__).resolve().parent
base='305648bc';source='1c3cfb9c97681415286236edd21a829e51b565eb';path='apps/artifacts/office-transform.mjs'
def git(*args):return subprocess.check_output(['git',*args],cwd=root)
def entries(ref):return dict((r.split(b'\t',1)[1].decode(),r.split(b'\t',1)[0].decode()) for r in git('ls-tree','-r','-z',ref).split(b'\0') if r)
before,after=entries(base),entries('HEAD');assert not(set(before)-set(after));changed=[p for p in before if before[p]!=after[p]];assert changed==[path],changed
old=git('show',base+':'+path);code=git('show',source+':'+path);assert code==(root/path).read_bytes()
args=lambda b:b.split(b'export function officeTransformArguments',1)[1].split(b'export function createLibreOfficeTransformer',1)[0]
assert args(old)==args(code)
groups=[]
for label in ['baseline','final']:
 r=json.loads((packet/(label+'.json')).read_text());assert r['sourceDiffBytes']==0 and r['ownedProcessGroupAbsent'] and not r['timedOut'] and not r['ownedGroupKillRequired'];groups.append(r['ownedPidAndPgid'])
 for stream in r['streams']:
  b=(root/stream['path']).read_bytes();assert len(b)==stream['bytes'] and hashlib.sha256(b).hexdigest()==stream['sha256']
 for name in ['original','styled']:
  receipt=json.loads((packet/(label+'-'+name+'.json')).read_text());data=(packet/(label+'-'+name+'.xlsx')).read_bytes();assert hashlib.sha256(data).hexdigest()==receipt['outputSha256'];assert receipt['nativeQualification'] is False and receipt['process']['groupAbsent'];groups.append(receipt['process']['pid'])
for group in groups:
 try:os.killpg(group,0)
 except ProcessLookupError:pass
 else:raise AssertionError('owned process group remains')
observations=[]
for file in ['original.xlsx','baseline-original.xlsx','final-original.xlsx']:
 w=load_workbook(packet/file,data_only=True);f=load_workbook(packet/file,data_only=False)
 observations.append({'file':file,'sha256':hashlib.sha256((packet/file).read_bytes()).hexdigest(),'cachedD2':w['Plan']['D2'].value,'cachedD3':w['Plan']['D3'].value,'cachedD4':w['Plan']['D4'].value,'D4Formula':f['Plan']['D4'].value,'NotesA2':f['Notes']['A2'].value,'NotesA2Type':f['Notes']['A2'].data_type});w.close();f.close()
assert [o['cachedD4'] for o in observations]==[0,0,2050]
report={'base':git('rev-parse',base).decode().strip(),'source':source,'originalEntries':len(before),'unchangedOriginalEntries':len(before)-1,'approvedOriginalChanges':changed,'sourceSha256':hashlib.sha256(code).hexdigest(),'officeArgumentFunctionByteIdentical':True,'allOtherOriginalEntriesIncludingPR53Unchanged':True,'retainedWorkbookObservations':observations,'ownedProcessGroupsFreshlyAbsent':groups,'dependenciesCreated':False,'nativeQualification':False}
(packet/'audit.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report))
